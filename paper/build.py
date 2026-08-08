# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from content1 import TITLE, SUBTITLE, PART1
from content2 import PART2
from content3 import PART3
from tables import TABLES

HERE = os.path.dirname(os.path.abspath(__file__))
IMGDIR = os.path.join(HERE, "img")
OUT = os.path.join(HERE, "国家5A级景区建设的减污降碳效应及作用机制.docx")

SONG = "宋体"
HEI = "黑体"
KAI = "楷体"
TNR = "Times New Roman"

BODY = PART1 + PART2 + PART3


# ----------------------------------------------------------------- helpers
def set_run_font(run, cn=SONG, en=TNR, size=None, bold=None, italic=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    rf.set(qn("w:cs"), en)


def set_style_font(style, cn, en, size, bold=False, color=None):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rPr = style.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    rf.set(qn("w:cs"), en)


def set_para_fmt(p, first_indent=None, before=0, after=0, line=1.5,
                 align=None, left_indent=None, hanging=None):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if left_indent is not None:
        pf.left_indent = Cm(left_indent)
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)


def cn_first_indent(p, chars=2):
    """按字符设置首行缩进（Word 的 firstLineChars）。"""
    pPr = p._element.get_or_add_pPr()
    ind = pPr.get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    ind.set(qn("w:firstLine"), "0")


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        succ = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
                "w:tcFitText", "w:vAlign", "w:hideMark", "w:headers",
                "w:tcPrChange")
        anchor = None
        for tag in succ:
            e = tcPr.find(qn(tag))
            if e is not None:
                anchor = e
                break
        if anchor is None:
            tcPr.append(borders)
        else:
            anchor.addprevious(borders)
    for name, spec in (("top", top), ("bottom", bottom),
                       ("left", left), ("right", right)):
        if spec is None:
            continue
        el = borders.find(qn("w:" + name))
        if el is None:
            el = OxmlElement("w:" + name)
            borders.append(el)
        el.set(qn("w:val"), spec[0])
        el.set(qn("w:sz"), str(spec[1]))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def clear_table_borders(table):
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + name)
        e.set(qn("w:val"), "none")
        e.set(qn("w:sz"), "0")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "auto")
        b.append(e)
    succ = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
            "w:tblCaption", "w:tblDescription", "w:tblPrChange")
    anchor = None
    for tag in succ:
        e = tblPr.find(qn(tag))
        if e is not None:
            anchor = e
            break
    if anchor is None:
        tblPr.append(b)
    else:
        anchor.addprevious(b)


def set_cell_text(cell, text, size=9.0, bold=False, align="c", cn=SONG):
    # 清除多余段落
    tc = cell._tc
    ps = tc.findall(qn("w:p"))
    for p in ps[1:]:
        tc.remove(p)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER,
                   "l": WD_ALIGN_PARAGRAPH.LEFT,
                   "r": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.space_before = Pt(1.5)
    pf.space_after = Pt(1.5)
    pf.line_spacing = 1.0
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.first_line_indent = Cm(0)
    if text:
        run = p.add_run(text)
        set_run_font(run, cn=cn, en=TNR, size=size, bold=bold)
    # 垂直居中
    tcPr = tc.get_or_add_tcPr()
    va = OxmlElement("w:vAlign")
    va.set(qn("w:val"), "center")
    tcPr.append(va)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


# ----------------------------------------------------------------- document
doc = Document()

# ---- 页面设置 ----
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(3.0)
sec.top_margin = Cm(2.8)
sec.bottom_margin = Cm(2.8)
CONTENT_W = 21.0 - 6.0  # 15.0 cm

# ---- 样式 ----
st = doc.styles["Normal"]
set_style_font(st, SONG, TNR, 12)
st.paragraph_format.line_spacing = 1.5
st.paragraph_format.space_after = Pt(0)

h_specs = {
    "Heading 1": (HEI, 14, WD_ALIGN_PARAGRAPH.CENTER, 14, 10),
    "Heading 2": (HEI, 12, WD_ALIGN_PARAGRAPH.LEFT, 10, 6),
    "Heading 3": (HEI, 12, WD_ALIGN_PARAGRAPH.LEFT, 8, 4),
}
for name, (cn, size, al, bef, aft) in h_specs.items():
    s = doc.styles[name]
    set_style_font(s, cn, TNR, size, bold=(name != "Heading 3"),
                   color=RGBColor(0, 0, 0))
    s.paragraph_format.alignment = al
    s.paragraph_format.space_before = Pt(bef)
    s.paragraph_format.space_after = Pt(aft)
    s.paragraph_format.line_spacing = 1.5
    s.paragraph_format.keep_with_next = True
    s.next_paragraph_style = doc.styles["Normal"]
    # 关闭"多级列表编号"，并确保 outlineLvl 存在（导航窗格依赖）
    pPr = s.element.get_or_add_pPr()
    for tag in ("w:numPr",):
        e = pPr.find(qn(tag))
        if e is not None:
            pPr.remove(e)
    lvl = pPr.find(qn("w:outlineLvl"))
    if lvl is None:
        lvl = OxmlElement("w:outlineLvl")
        pPr.append(lvl)
    lvl.set(qn("w:val"), str(int(name[-1]) - 1))


def add_p(text, style=None):
    return doc.add_paragraph(text, style=style)


# ---- 标题 ----
p = add_p("")
set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, line=1.5)
r = p.add_run(TITLE)
set_run_font(r, cn=HEI, en=TNR, size=18, bold=True)

p = add_p("")
set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=16, line=1.5)
r = p.add_run(SUBTITLE)
set_run_font(r, cn=HEI, en=TNR, size=14, bold=False)


# ---- TOC 域 ----
def add_toc():
    p = doc.add_paragraph()
    set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=6, line=1.5)
    r = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    r._r.append(fld)
    r2 = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-2" \\h \\z \\u '
    r2._r.append(it)
    r3 = p.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    r4 = p.add_run("【目录将在打开文档后自动生成；如未显示，请在此处右键选择“更新域”或按 F9】")
    set_run_font(r4, cn=SONG, en=TNR, size=10.5)
    r5 = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5._r.append(end)


def add_table(tid):
    spec = TABLES[tid]
    rows = spec["rows"]
    nr, nc = len(rows), len(rows[0])
    widths = spec["w"]
    tot = sum(widths)
    widths_cm = [CONTENT_W * w / tot for w in widths]

    t = doc.add_table(rows=nr, cols=nc)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    clear_table_borders(t)

    skip = set()
    for (r1, c1, r2, c2) in spec.get("merges", []):
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                if (rr, cc) != (r1, c1):
                    skip.add((rr, cc))
        t.cell(r1, c1).merge(t.cell(r2, c2))

    hdr = spec["hdr"]
    align = spec["align"]
    for i in range(nr):
        for j in range(nc):
            if (i, j) in skip:
                continue
            txt = rows[i][j]
            is_hdr = i < hdr
            a = "c" if is_hdr else align[j]
            set_cell_text(t.cell(i, j), txt, size=9.0, bold=is_hdr, align=a,
                          cn=(HEI if is_hdr else SONG))

    for i in range(nr):
        for j in range(nc):
            t.cell(i, j).width = Cm(widths_cm[j])

    # 三线表边框
    for j in range(nc):
        set_cell_borders(t.cell(0, j), top=("single", 12))
        set_cell_borders(t.cell(hdr - 1, j), bottom=("single", 6))
        set_cell_borders(t.cell(nr - 1, j), bottom=("single", 12))
    for i in range(hdr):
        set_repeat_header(t.rows[i])
    return t


def add_note(text, size=9.0, indent=True):
    p = doc.add_paragraph()
    set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=3, after=8, line=1.25)
    p.paragraph_format.left_indent = Cm(0)
    r = p.add_run(text)
    set_run_font(r, cn=SONG, en=TNR, size=size)


def add_eq(text, num):
    p = doc.add_paragraph()
    set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=8, line=1.5)
    p.paragraph_format.first_line_indent = Cm(0)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Cm(CONTENT_W / 2.0), WD_TAB_ALIGNMENT.CENTER)
    tabs.add_tab_stop(Cm(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("\t" + text + "\t" + num)
    set_run_font(r, cn=SONG, en=TNR, size=12, italic=False)


for item in BODY:
    t = item["t"]
    x = item.get("x", "")

    if t == "h1":
        p = doc.add_paragraph(x, style="Heading 1")
    elif t == "h2":
        p = doc.add_paragraph(x, style="Heading 2")
    elif t == "h3":
        p = doc.add_paragraph(x, style="Heading 3")

    elif t == "p":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, line=1.5)
        cn_first_indent(p, 2)
        r = p.add_run(x)
        set_run_font(r, cn=SONG, en=TNR, size=12)

    elif t == "hyp":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=4, after=4, line=1.5)
        cn_first_indent(p, 2)
        r = p.add_run(x)
        set_run_font(r, cn=KAI, en=TNR, size=12, bold=True)

    elif t == "eq":
        add_eq(x, item["n"])

    elif t == "tabcap":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=4, line=1.25)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(x)
        set_run_font(r, cn=HEI, en=TNR, size=10.5, bold=True)

    elif t == "table":
        add_table(item["id"])

    elif t == "note":
        add_note(x)

    elif t == "note_c":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=8, line=1.25)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(x)
        set_run_font(r, cn=SONG, en=TNR, size=9)

    elif t == "img":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2, line=1.0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(os.path.join(IMGDIR, x), width=Cm(item["w"]))

    elif t == "img2":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=2, line=1.0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(os.path.join(IMGDIR, x[0]), width=Cm(item["w"]))
        run2 = p.add_run("  ")
        set_run_font(run2, size=9)
        run3 = p.add_run()
        run3.add_picture(os.path.join(IMGDIR, x[1]), width=Cm(item["w"]))

    elif t == "figcap":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10, line=1.25)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(x)
        set_run_font(r, cn=HEI, en=TNR, size=10.5, bold=True)

    elif t == "ref":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=2, line=1.35)
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.9)
        r = p.add_run(x)
        set_run_font(r, cn=SONG, en=TNR, size=10.5)

    elif t == "abs_h":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3, line=1.4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(x)
        set_run_font(r, cn=HEI, en=TNR, size=10.5, bold=True)

    elif t == "abs_p":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, line=1.4)
        cn_first_indent(p, 2)
        r = p.add_run(x)
        set_run_font(r, cn=SONG, en=TNR, size=10.5)

    elif t == "abs_en":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, line=1.4)
        p.paragraph_format.first_line_indent = Cm(0.7)
        r = p.add_run(x)
        set_run_font(r, cn=TNR, en=TNR, size=10.5)

    elif t == "abs_kw":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=3, line=1.4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(x)
        set_run_font(r, cn=HEI, en=TNR, size=10.5)

    elif t == "abs_kw_en":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=3, line=1.4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(x)
        set_run_font(r, cn=TNR, en=TNR, size=10.5, bold=True)

    elif t == "en_title":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4, line=1.4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(x)
        set_run_font(r, cn=TNR, en=TNR, size=14, bold=True)

    elif t == "en_sub":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=10, line=1.4)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run("——" + x)
        set_run_font(r, cn=TNR, en=TNR, size=12)

    elif t == "toc":
        p = doc.add_paragraph()
        set_para_fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, line=1.5)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run("目　　录")
        set_run_font(r, cn=HEI, en=TNR, size=14, bold=True)
        add_toc()

    elif t == "pagebreak":
        doc.add_page_break()

    else:
        raise ValueError("unknown item type: " + t)


# ---- 页脚页码 ----
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
r2 = fp.add_run()
it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
r2._r.append(it)
r3 = fp.add_run()
f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r3._r.append(f2)
for rr in fp.runs:
    set_run_font(rr, cn=SONG, en=TNR, size=10.5)

# ---- 打开时更新域（使目录自动生成） ----
settings = doc.settings.element
zoom = settings.find(qn("w:zoom"))
if zoom is not None:
    zoom.set(qn("w:percent"), "100")
uf = settings.find(qn("w:updateFields"))
if uf is None:
    uf = OxmlElement("w:updateFields")
    compat = settings.find(qn("w:compat"))
    if compat is not None:
        compat.addprevious(uf)
    else:
        settings.append(uf)
uf.set(qn("w:val"), "true")

doc.save(OUT)
print("saved:", OUT)
